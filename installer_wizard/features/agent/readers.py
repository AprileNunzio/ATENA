import asyncio
import shutil
from pathlib import Path

OFFICE = {".docx", ".xlsx", ".pptx"}


def _docx(p: Path) -> str:
    from docx import Document
    doc = Document(str(p))
    rows = [para.text for para in doc.paragraphs]
    for table in doc.tables:
        rows += [" | ".join(c.text.strip() for c in r.cells) for r in table.rows]
    return "\n".join(rows)


def _xlsx(p: Path) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(str(p), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"## {ws.title}")
        for n, row in enumerate(ws.iter_rows(values_only=True)):
            if n >= 500:
                break
            out.append(" | ".join("" if v is None else str(v) for v in row))
    wb.close()
    return "\n".join(out)


def _pptx(p: Path) -> str:
    from pptx import Presentation
    out = []
    for n, slide in enumerate(Presentation(str(p)).slides, 1):
        out.append(f"## Diapositiva {n}")
        out += [s.text_frame.text for s in slide.shapes if s.has_text_frame]
    return "\n".join(out)


async def _pdf(p: Path) -> str:
    if not shutil.which("pdftotext"):
        raise ValueError("per leggere i PDF serve pdftotext (pacchetto poppler-utils)")
    proc = await asyncio.create_subprocess_exec("pdftotext", "-layout", str(p), "-", stdout=asyncio.subprocess.PIPE,
                                                stderr=asyncio.subprocess.PIPE)
    out, err = await asyncio.wait_for(proc.communicate(), 60)
    if proc.returncode:
        raise ValueError(f"PDF non leggibile: {err.decode('utf-8', 'replace')[:200]}")
    return out.decode("utf-8", "replace")


async def text_of(p: Path) -> str:
    suffix = p.suffix.lower()
    if suffix == ".pdf":
        return await _pdf(p)
    if suffix in OFFICE:
        try:
            reader = {".docx": _docx, ".xlsx": _xlsx, ".pptx": _pptx}[suffix]
            return await asyncio.to_thread(reader, p)
        except ImportError as exc:
            raise ValueError(f"manca la libreria per leggere i file {suffix}: {exc.name}") from exc
    return p.read_text(encoding="utf-8", errors="replace")
